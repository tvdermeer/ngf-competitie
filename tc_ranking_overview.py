#!/usr/bin/env python3
"""Build the TC Ranking pages from the ranking workbook.

    python tc_ranking_overview.py

Reads "2027 TC Ranking 14092026.xlsx" (sheet "Ranking") and writes two
self-contained HTML pages:

    2027 TC Ranking overzicht.html   read-only kanban of the current ranking
    2027 Indeling.html               drag & drop board to build the 2027 teams

Every player becomes a card with their name, "aantal holes" preference and
average SD. Cards are grouped into lanes named after the team the player was
in during 2026; players without a 2026 team sit in one shared "didn't play"
lane. The indeling page remembers its layout in the browser and can export or
import it as JSON.

With --encrypt the embedded data is encrypted (PBKDF2-SHA256 + AES-256-GCM) and
the pages ask for a password before showing anything. Because GitHub Pages is
public, this is the only way to keep the names out of reach: a plain JavaScript
password gate would leave the data readable in the page source.

    python tc_ranking_overview.py --encrypt            (prompts for a password)
    python tc_ranking_overview.py --encrypt --docs     (output for GitHub Pages)

Encryption needs Node.js on the build machine (used only at build time).
"""

from __future__ import annotations

import argparse
import csv
import getpass
import json
import os
import re
import shutil
import subprocess
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent
XLSX = ROOT / "2027 TC Ranking 14092026.xlsx"
SUMMARY_CSV = ROOT / "tc_ranking_summary.csv"
ENCRYPTOR = ROOT / "encrypt_payload.js"
DOCS_DIR = ROOT / "docs"
PBKDF2_ITERATIONS = 600000

# Filenames differ between the plain build and the GitHub Pages (docs) build so
# that the published entry point is index.html.
OVERVIEW_NAME = "2027 TC Ranking overzicht.html"
INDELING_NAME = "2027 Indeling.html"

SHEET = "Ranking"
NO_TEAM_LABEL = "Niet gespeeld 2026"

CATEGORY_ORDER = [
    "Heren",
    "Heren 50+",
    "Dames",
    "Dames 50+",
    "C-Mixed",
    NO_TEAM_LABEL,
]


def natural_key(value: str):
    return [int(part) if part.isdigit() else part.lower()
            for part in re.split(r"(\d+)", value)]


def category_of(team: str) -> str:
    if team == NO_TEAM_LABEL:
        return NO_TEAM_LABEL
    if team.startswith("Heren 50+"):
        return "Heren 50+"
    if team.startswith("Heren"):
        return "Heren"
    if team.startswith("Dames 50+"):
        return "Dames 50+"
    if team.startswith("Dames"):
        return "Dames"
    if team.startswith("C-Mixed"):
        return "C-Mixed"
    return team


def lane_sort_key(team: str):
    cat = category_of(team)
    return (CATEGORY_ORDER.index(cat) if cat in CATEGORY_ORDER else 99,
            natural_key(team))


def parse_holes(raw) -> str:
    if raw is None:
        return ""
    text = str(raw).strip()
    if not text:
        return ""
    if "geen" in text.lower():
        return "Geen voorkeur"
    parts = [p.strip() for p in text.split(";") if p.strip()]
    return "/".join(parts)


def parse_sd(raw):
    if raw is None:
        return None
    text = str(raw).strip().replace(",", ".")
    if not text:
        return None
    try:
        return round(float(text), 2)
    except ValueError:
        return None


def find_header(ws):
    for idx, row in enumerate(ws.iter_rows(min_row=1, max_row=30, values_only=True), 1):
        if row and row[0] == "Lidcode":
            return idx, list(row)
    raise SystemExit("Kon de kopregel (Lidcode) niet vinden in blad 'Ranking'.")


def load_ratings():
    """Rating (Rt) per player from the NGF summary, keyed by lidcode."""
    if not SUMMARY_CSV.exists():
        print(f"Let op: {SUMMARY_CSV.name} niet gevonden; rating blijft leeg.")
        return {}
    ratings = {}
    with SUMMARY_CSV.open(encoding="utf-8-sig", newline="") as fh:
        for row in csv.DictReader(fh, delimiter=";"):
            code = (row.get("Lidcode") or "").strip()
            raw = (row.get("Rating (Rt)") or "").strip().replace(",", ".")
            if not code or not raw:
                continue
            try:
                ratings[code] = round(float(raw), 2)
            except ValueError:
                continue
    return ratings


def load_players(ratings):
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb[SHEET]
    header_idx, header = find_header(ws)

    col = {name: header.index(name) for name in header if name}
    idx_name = col["Naam"]
    idx_holes = col["Aantal holes"]
    idx_team = col["Team 2026"]
    idx_sd = col["Gem. SD"]
    idx_code = col["Lidcode"]

    players: list[dict] = []
    seen = set()
    for row in ws.iter_rows(min_row=header_idx + 1, values_only=True):
        if not row or not row[idx_name]:
            continue
        name = str(row[idx_name]).strip()
        team = (str(row[idx_team]).strip() if row[idx_team] is not None else "")
        team = team or NO_TEAM_LABEL
        sd = parse_sd(row[idx_sd])
        code = (str(row[idx_code]).strip() if row[idx_code] is not None else "")
        key = (code, name, team, sd)
        if key in seen:
            continue
        seen.add(key)
        players.append({
            "id": code,
            "code": code,
            "name": name,
            "holes": parse_holes(row[idx_holes]),
            "team": team,
            "sd": sd,
            "rating": ratings.get(code),
        })
    return players


def build_lanes(players):
    groups: dict[str, list] = {}
    for player in players:
        groups.setdefault(player["team"], []).append(player)

    lanes = []
    for team in sorted(groups, key=lane_sort_key):
        cards = sorted(groups[team],
                       key=lambda p: (p["sd"] is None, p["sd"] if p["sd"] is not None else 0))
        lanes.append({
            "team": team,
            "category": category_of(team),
            "noTeam": team == NO_TEAM_LABEL,
            "players": cards,
        })
    return lanes


BASE_CSS = r"""
  :root {
    --bg: #eef1f5;
    --panel: #ffffff;
    --ink: #1c2430;
    --muted: #66707c;
    --line: #dfe4ea;
    --shadow: 0 1px 2px rgba(20,30,45,.08), 0 8px 24px rgba(20,30,45,.06);
  }
  * { box-sizing: border-box; }
  html, body { height: 100%; }
  body {
    margin: 0;
    font-family: "Segoe UI", system-ui, -apple-system, Roboto, Helvetica, Arial, sans-serif;
    color: var(--ink);
    background: radial-gradient(1200px 600px at 15% -10%, #ffffff 0%, var(--bg) 55%);
    display: flex;
    flex-direction: column;
    height: 100vh;
    overflow: hidden;
  }
  header.top {
    padding: 16px 24px 13px;
    display: flex;
    flex-wrap: wrap;
    gap: 12px 22px;
    align-items: center;
    border-bottom: 1px solid var(--line);
    background: rgba(255,255,255,.8);
    backdrop-filter: blur(8px);
  }
  .title h1 { margin: 0; font-size: 19px; letter-spacing: .2px; }
  .title p { margin: 3px 0 0; font-size: 12.5px; color: var(--muted); }
  .stats { display: flex; gap: 18px; flex-wrap: wrap; }
  .stat { text-align: center; }
  .stat b { display: block; font-size: 20px; line-height: 1; }
  .stat span { font-size: 11px; text-transform: uppercase; letter-spacing: .7px; color: var(--muted); }
  .spacer { margin-left: auto; }
  .tools { display: flex; gap: 10px; align-items: center; flex-wrap: wrap; }
  input[type=search], select {
    font: inherit;
    font-size: 13px;
    padding: 7px 11px;
    border: 1px solid var(--line);
    border-radius: 9px;
    background: #fff;
    color: var(--ink);
    outline: none;
  }
  input[type=search]:focus, select:focus { border-color: #6b8cff; box-shadow: 0 0 0 3px rgba(107,140,255,.18); }
  .btn {
    font: inherit;
    font-size: 13px;
    font-weight: 600;
    padding: 7px 13px;
    border: 1px solid var(--line);
    border-radius: 9px;
    background: #fff;
    color: var(--ink);
    cursor: pointer;
    white-space: nowrap;
    text-decoration: none;
  }
  .btn:hover { border-color: #b9c2cd; background: #f7f9fc; }
  .btn.primary { background: #3b6fd4; border-color: #3b6fd4; color: #fff; }
  .btn.primary:hover { background: #335fbb; }
  .btn.ghost { color: var(--muted); }
  main.board {
    flex: 1;
    display: flex;
    gap: 16px;
    padding: 18px 24px 24px;
    overflow-x: auto;
    overflow-y: hidden;
    align-items: stretch;
  }
  .lane {
    flex: 0 0 258px;
    width: 258px;
    display: flex;
    flex-direction: column;
    background: var(--panel);
    border-radius: 14px;
    box-shadow: var(--shadow);
    border: 1px solid var(--line);
    overflow: hidden;
    max-height: 100%;
  }
  .lane.hidden { display: none; }
  .lane > header {
    padding: 11px 12px;
    border-bottom: 1px solid var(--line);
    display: flex;
    align-items: center;
    gap: 8px;
    border-top: 4px solid var(--accent, #6b8cff);
    background: linear-gradient(180deg, color-mix(in srgb, var(--accent, #6b8cff) 9%, #fff), #fff);
  }
  .lane > header h2 {
    margin: 0;
    font-size: 13.5px;
    font-weight: 700;
    flex: 1;
    letter-spacing: .1px;
    outline: none;
    border-radius: 5px;
    padding: 1px 3px;
    min-width: 0;
  }
  .lane > header h2[contenteditable=true]:focus { background: #fff; box-shadow: 0 0 0 2px rgba(107,140,255,.35); }
  .lane > header .count {
    font-size: 11.5px;
    font-weight: 700;
    color: var(--accent, #6b8cff);
    background: color-mix(in srgb, var(--accent, #6b8cff) 14%, #fff);
    border: 1px solid color-mix(in srgb, var(--accent, #6b8cff) 30%, #fff);
    border-radius: 999px;
    padding: 1px 9px;
    flex: 0 0 auto;
  }
  .lane.noTeam { --accent: #8b95a3; }
  .cards {
    padding: 10px;
    overflow-y: auto;
    display: flex;
    flex-direction: column;
    gap: 8px;
    min-height: 54px;
    scrollbar-width: thin;
  }
  .cards::-webkit-scrollbar { width: 9px; }
  .cards::-webkit-scrollbar-thumb { background: #cfd6df; border-radius: 9px; border: 3px solid #fff; }
  .card {
    border: 1px solid var(--line);
    border-left: 4px solid var(--hole, #b9c2cd);
    border-radius: 10px;
    padding: 9px 10px 10px;
    background: #fff;
    transition: transform .08s ease, box-shadow .12s ease, opacity .12s ease;
  }
  .card:hover { transform: translateY(-1px); box-shadow: 0 6px 16px rgba(20,30,45,.10); }
  .card .name { font-size: 13.5px; font-weight: 600; line-height: 1.25; word-break: break-word; }
  .card .meta { display: flex; align-items: center; gap: 6px; margin-top: 7px; flex-wrap: wrap; }
  .badge {
    font-size: 11px;
    font-weight: 600;
    padding: 2px 8px;
    border-radius: 999px;
    white-space: nowrap;
  }
  .holes { background: var(--hole-bg, #eef1f5); color: var(--hole-fg, #4a5563); border: 1px solid var(--hole-bd, #dde2e8); }
  .sd { background: var(--sd-bg); color: var(--sd-fg); border: 1px solid var(--sd-bd); }
  .sd.none { background: #f2f4f7; color: #98a1ac; border-color: #e6e9ee; }
  .empty { color: var(--muted); font-size: 13px; padding: 20px 12px; text-align: center; }
  .legend { display: flex; gap: 12px; align-items: center; font-size: 11.5px; color: var(--muted); flex-wrap: wrap; }
  .legend i { display: inline-block; width: 11px; height: 11px; border-radius: 3px; margin-right: 4px; vertical-align: -1px; }
  @media print {
    body { height: auto; overflow: visible; background: #fff; }
    header.top { position: static; }
    main.board { flex-wrap: wrap; overflow: visible; }
    .lane { max-height: none; break-inside: avoid; }
    .cards { overflow: visible; }
    header.top .tools, header.top .only-screen { display: none; }
    .lock { display: none !important; }
  }
  .locked { overflow: hidden; }
  body.locked main, body.locked header.top { visibility: hidden; }
  .lock {
    position: fixed;
    inset: 0;
    z-index: 50;
    display: flex;
    align-items: center;
    justify-content: center;
    padding: 20px;
    background: radial-gradient(900px 500px at 50% 0%, #ffffff 0%, #e7ebf1 70%);
  }
  .lock[hidden] { display: none; }
  .lock-card {
    width: min(380px, 100%);
    background: #fff;
    border: 1px solid var(--line);
    border-radius: 16px;
    box-shadow: var(--shadow);
    padding: 26px 24px 22px;
    text-align: center;
  }
  .lock-card .lock-icon { font-size: 30px; line-height: 1; }
  .lock-card h2 { margin: 10px 0 4px; font-size: 18px; }
  .lock-card p { margin: 0 0 16px; font-size: 13px; color: var(--muted); }
  .lock-card input {
    width: 100%;
    font: inherit;
    font-size: 14px;
    padding: 10px 12px;
    border: 1px solid var(--line);
    border-radius: 9px;
    outline: none;
    margin-bottom: 10px;
  }
  .lock-card input:focus { border-color: #6b8cff; box-shadow: 0 0 0 3px rgba(107,140,255,.18); }
  .lock-card .btn { width: 100%; padding: 10px; }
  .lock-error { color: #e0483b; font-size: 12.5px; min-height: 16px; margin: 10px 0 0; }
  .info-group { display: flex; gap: 6px; align-items: center; flex-wrap: wrap; }
  .info-btn {
    display: inline-flex;
    align-items: center;
    gap: 6px;
    font: inherit;
    font-size: 12px;
    font-weight: 600;
    color: var(--ink);
    background: #fff;
    border: 1px solid var(--line);
    border-radius: 999px;
    padding: 4px 11px 4px 5px;
    cursor: pointer;
  }
  .info-btn:hover { border-color: #6b8cff; background: #f4f7ff; }
  .info-icon {
    width: 17px;
    height: 17px;
    border-radius: 50%;
    background: #3b6fd4;
    color: #fff;
    font-family: Georgia, "Times New Roman", serif;
    font-style: italic;
    font-size: 12px;
    line-height: 1;
    display: inline-flex;
    align-items: center;
    justify-content: center;
  }
  .modal {
    position: fixed;
    inset: 0;
    z-index: 60;
    display: flex;
    align-items: center;
    justify-content: center;
    padding: 20px;
    background: rgba(20,30,45,.45);
  }
  .modal[hidden] { display: none; }
  .modal-card {
    width: min(440px, 100%);
    background: #fff;
    border: 1px solid var(--line);
    border-radius: 14px;
    box-shadow: var(--shadow);
    padding: 22px 22px 18px;
  }
  .modal-card h2 { margin: 0 0 8px; font-size: 17px; }
  .modal-card p { margin: 0 0 18px; font-size: 13.5px; line-height: 1.55; color: #3a4553; }
  .modal-card .btn { width: 100%; }
  .lock-card.shake { animation: shake .3s; }
  @keyframes shake {
    0%, 100% { transform: translateX(0); }
    25% { transform: translateX(-6px); }
    75% { transform: translateX(6px); }
  }
"""

INDELING_CSS = r"""
  .lane.drop-target { box-shadow: 0 0 0 2px var(--accent, #6b8cff), var(--shadow); }
  .card { cursor: grab; }
  .card:active { cursor: grabbing; }
  .card.dragging { opacity: .35; }
  .card.dimmed { opacity: .28; }
  .placeholder {
    border: 2px dashed var(--accent, #6b8cff);
    border-radius: 10px;
    background: color-mix(in srgb, var(--accent, #6b8cff) 7%, #fff);
    min-height: 46px;
  }
  .lane .del {
    border: none;
    background: transparent;
    color: #b3bcc7;
    font-size: 16px;
    line-height: 1;
    cursor: pointer;
    padding: 2px 4px;
    border-radius: 6px;
    flex: 0 0 auto;
  }
  .lane .del:hover { color: #e0483b; background: #fdecea; }
  .saved {
    font-size: 12px;
    color: #2f9e7d;
    opacity: 0;
    transition: opacity .3s ease;
    min-width: 62px;
  }
  .saved.show { opacity: 1; }
  .hint { font-size: 12px; color: var(--muted); }
"""

UNLOCK_HTML = r"""
<div class="lock" id="lock" hidden>
  <form class="lock-card" id="lock-form">
    <div class="lock-icon">&#128274;</div>
    <h2>Beveiligde pagina</h2>
    <p>Deze pagina bevat persoonsgegevens. Voer het wachtwoord in om verder te gaan.</p>
    <input type="password" id="lock-pw" placeholder="Wachtwoord" autocomplete="current-password" required>
    <button class="btn primary" type="submit">Ontgrendelen</button>
    <p class="lock-error" id="lock-error"></p>
  </form>
</div>
"""

INFO_HTML = r"""
<div class="modal" id="info-modal" hidden>
  <div class="modal-card" role="dialog" aria-modal="true" aria-labelledby="info-title">
    <h2 id="info-title"></h2>
    <p id="info-body"></p>
    <button class="btn primary" id="info-close" type="button">Sluiten</button>
  </div>
</div>
"""

INFO_JS = r"""
const INFO_TEXTS = {
  sd: {
    title: 'SD',
    body: 'SD is de gemiddelde hcp (het dagresultaat) in de Q-kaarten (qualifying-wedstrijden) over het afgelopen jaar.'
  },
  rating: {
    title: 'Rating',
    body: 'Rating is een scoresysteem voor de prestatie tijdens de afgelopen competitie. Je krijgt punten voor gewonnen wedstrijden en aftrek voor verloren wedstrijden. Op basis van het verschil in hcp van de tegenstander krijgt de gewonnen of verloren wedstrijd een gewicht.'
  }
};

(function () {
  const modal = document.getElementById('info-modal');
  if (!modal) return;
  const titleEl = document.getElementById('info-title');
  const bodyEl = document.getElementById('info-body');

  function open(key) {
    const info = INFO_TEXTS[key];
    if (!info) return;
    titleEl.textContent = info.title;
    bodyEl.textContent = info.body;
    modal.hidden = false;
  }
  function close() { modal.hidden = true; }

  Array.prototype.forEach.call(document.querySelectorAll('.info-btn'), function (btn) {
    btn.addEventListener('click', function () { open(btn.dataset.info); });
  });
  document.getElementById('info-close').addEventListener('click', close);
  modal.addEventListener('click', function (ev) { if (ev.target === modal) close(); });
  document.addEventListener('keydown', function (ev) {
    if (ev.key === 'Escape' && !modal.hidden) close();
  });
})();
"""

UNLOCK_JS = r"""
const SESSION_KEY = 'ngf-tc-unlocked-v1';

function b64ToBuf(b64) {
  const bin = atob(b64);
  const bytes = new Uint8Array(bin.length);
  for (let i = 0; i < bin.length; i++) bytes[i] = bin.charCodeAt(i);
  return bytes;
}

async function deriveKey(password, salt, iterations) {
  const base = await crypto.subtle.importKey(
    'raw', new TextEncoder().encode(password), { name: 'PBKDF2' }, false, ['deriveKey']);
  return crypto.subtle.deriveKey(
    { name: 'PBKDF2', salt: salt, iterations: iterations, hash: 'SHA-256' },
    base, { name: 'AES-GCM', length: 256 }, false, ['decrypt']);
}

async function decryptPayload(enc, password) {
  const key = await deriveKey(password, b64ToBuf(enc.salt), enc.iterations);
  const plain = await crypto.subtle.decrypt({ name: 'AES-GCM', iv: b64ToBuf(enc.iv) }, key, b64ToBuf(enc.ct));
  return JSON.parse(new TextDecoder().decode(plain));
}

function showLock() {
  document.body.classList.add('locked');
  document.getElementById('lock').hidden = false;
}

function hideLock() {
  document.body.classList.remove('locked');
  document.getElementById('lock').hidden = true;
}

async function unlockFlow(payload, onReady) {
  if (!payload.enc) { onReady(payload.plain); return; }

  const cached = sessionStorage.getItem(SESSION_KEY);
  if (cached) {
    try { onReady(JSON.parse(cached)); return; } catch (e) { sessionStorage.removeItem(SESSION_KEY); }
  }

  if (!window.crypto || !crypto.subtle) {
    showLock();
    document.getElementById('lock-error').textContent =
      'Decoderen lukt niet in deze browser of context. Gebruik https of een moderne browser.';
    return;
  }

  showLock();
  const form = document.getElementById('lock-form');
  const card = form.querySelector('.lock-card') || form;
  const pw = document.getElementById('lock-pw');
  const error = document.getElementById('lock-error');
  pw.focus();

  form.addEventListener('submit', async function (ev) {
    ev.preventDefault();
    error.textContent = '';
    const btn = form.querySelector('button[type=submit]');
    btn.disabled = true;
    try {
      const data = await decryptPayload(payload.enc, pw.value);
      try { sessionStorage.setItem(SESSION_KEY, JSON.stringify(data)); } catch (e) { /* ignore */ }
      hideLock();
      onReady(data);
    } catch (e) {
      error.textContent = 'Onjuist wachtwoord. Probeer het opnieuw.';
      card.classList.remove('shake');
      void card.offsetWidth;
      card.classList.add('shake');
      pw.select();
    } finally {
      btn.disabled = false;
    }
  });
}
"""

OVERVIEW_TEMPLATE = r"""<!DOCTYPE html>
<html lang="nl">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<meta name="robots" content="noindex, nofollow">
<title>2027 TC Ranking &middot; overzicht</title>
<style>__CSS__</style>
</head>
<body>
<header class="top">
  <div class="title">
    <h1>2027 TC Ranking &middot; overzicht</h1>
    <p>Kaarten gegroepeerd per team in 2026 &middot; naam, aantal holes, gemiddeld SD en rating 2026</p>
  </div>
  <div class="info-group">
    <button type="button" class="info-btn" data-info="sd"><span class="info-icon">i</span>SD</button>
    <button type="button" class="info-btn" data-info="rating"><span class="info-icon">i</span>Rating</button>
  </div>
  <div class="legend" id="legend"></div>
  <div class="spacer"></div>
  <div class="stats" id="stats"></div>
  <div class="tools">
    <input type="search" id="q" placeholder="Zoek naam&hellip;" autocomplete="off">
    <select id="sort">
      <option value="sd-asc">SD laag &rarr; hoog</option>
      <option value="sd-desc">SD hoog &rarr; laag</option>
      <option value="name">Naam A &rarr; Z</option>
    </select>
    <a class="btn primary" href="__INDELING_HREF__">Naar 2027 indeling &rarr;</a>
  </div>
</header>
<main class="board" id="board"></main>
__UNLOCK_HTML__
__INFO_HTML__
<script id="payload" type="application/json">__DATA__</script>
<script>__UNLOCK_JS__</script>
<script>__INFO_JS__</script>
<script>
const PAYLOAD = JSON.parse(document.getElementById('payload').textContent);
const board = document.getElementById('board');
const q = document.getElementById('q');
const sortSel = document.getElementById('sort');

function start(DATA) {

const CAT_COLORS = {
  'Heren': '#3b6fd4',
  'Heren 50+': '#c98a1a',
  'Dames': '#d1478a',
  'Dames 50+': '#9b5cc4',
  'C-Mixed': '#2f9e7d',
  'Niet gespeeld 2026': '#8b95a3'
};

const HOLE_COLORS = {
  '36': ['#e3edff', '#1f4fa8', '#bcd3ff'],
  '27': ['#e2f5ee', '#0f7a5a', '#bce8d8'],
  '18': ['#f1e8fb', '#6b33a8', '#ddc9f5'],
  '18/27': ['#e8eefc', '#3d5bbf', '#cbd7f7'],
  '27/36': ['#e0f4f8', '#10708a', '#bfe6ef'],
  'Geen voorkeur': ['#f2f4f7', '#6b7480', '#e4e8ee'],
  '': ['#f2f4f7', '#6b7480', '#e4e8ee']
};

function sdScale(sd) {
  const min = 3, max = 36;
  const t = Math.min(1, Math.max(0, (sd - min) / (max - min)));
  const hue = (1 - t) * 128;
  return {
    bg: 'hsl(' + hue + ' 72% 93%)',
    fg: 'hsl(' + hue + ' 55% 30%)',
    bd: 'hsl(' + hue + ' 60% 82%)'
  };
}

function fmtSd(sd) {
  if (sd === null || sd === undefined) return null;
  return sd.toFixed(1).replace('.', ',');
}

function fmtRating(rating) {
  if (rating === null || rating === undefined) return null;
  const v = Math.round(rating * 100) / 100;
  let body = Math.abs(v).toFixed(2).replace(/0+$/, '').replace(/\.$/, '');
  if (body === '') body = '0';
  const sign = v > 0 ? '+' : (v < 0 ? '-' : '');
  return sign + body.replace('.', ',');
}

function ratingColors(rating) {
  const v = Math.round(rating * 100) / 100;
  if (v === 0) return { bg: '#f2f4f7', fg: '#6b7480', bd: '#e4e8ee' };
  const hue = v > 0 ? 140 : 0;
  const mag = Math.min(1, Math.abs(v) / 7);
  const bgL = 94 - mag * 9;
  return {
    bg: 'hsl(' + hue + ' 70% ' + bgL + '%)',
    fg: 'hsl(' + hue + ' 62% ' + (30 - mag * 4) + '%)',
    bd: 'hsl(' + hue + ' 58% ' + (bgL - 9) + '%)'
  };
}

function makeCard(p) {
  const card = document.createElement('div');
  card.className = 'card';
  card.dataset.name = p.name.toLowerCase();
  const hc = HOLE_COLORS[p.holes] || HOLE_COLORS[''];
  card.style.setProperty('--hole', hc[1]);
  card.style.setProperty('--hole-bg', hc[0]);
  card.style.setProperty('--hole-fg', hc[1]);
  card.style.setProperty('--hole-bd', hc[2]);

  const name = document.createElement('div');
  name.className = 'name';
  name.textContent = p.name;
  card.appendChild(name);

  const meta = document.createElement('div');
  meta.className = 'meta';

  const holes = document.createElement('span');
  holes.className = 'badge holes';
  holes.textContent = p.holes ? p.holes + ' holes' : 'holes onbekend';
  meta.appendChild(holes);

  const sd = document.createElement('span');
  const val = fmtSd(p.sd);
  if (val === null) {
    sd.className = 'badge sd none';
    sd.textContent = 'SD n.v.t.';
  } else {
    sd.className = 'badge sd';
    const c = sdScale(p.sd);
    sd.style.setProperty('--sd-bg', c.bg);
    sd.style.setProperty('--sd-fg', c.fg);
    sd.style.setProperty('--sd-bd', c.bd);
    sd.textContent = 'SD ' + val;
    sd.title = 'Gemiddeld SD ' + p.sd;
  }
  meta.appendChild(sd);

  const rt = document.createElement('span');
  const rval = fmtRating(p.rating);
  if (rval === null) {
    rt.className = 'badge sd none';
    rt.textContent = 'Rating n.v.t.';
  } else {
    rt.className = 'badge sd';
    const rc = ratingColors(p.rating);
    rt.style.setProperty('--sd-bg', rc.bg);
    rt.style.setProperty('--sd-fg', rc.fg);
    rt.style.setProperty('--sd-bd', rc.bd);
    rt.textContent = 'Rating ' + rval;
    rt.title = 'Rating 2026 (Rt) ' + p.rating;
  }
  meta.appendChild(rt);

  card.appendChild(meta);
  return card;
}

function render() {
  board.innerHTML = '';
  const term = q.value.trim().toLowerCase();
  const mode = sortSel.value;

  for (const lane of DATA.lanes) {
    const section = document.createElement('section');
    section.className = 'lane' + (lane.noTeam ? ' noTeam' : '');
    section.style.setProperty('--accent', CAT_COLORS[lane.category] || '#6b8cff');

    const players = lane.players.filter(p => !term || p.name.toLowerCase().includes(term));

    const header = document.createElement('header');
    const h2 = document.createElement('h2');
    h2.textContent = lane.team;
    header.appendChild(h2);
    const count = document.createElement('span');
    count.className = 'count';
    count.textContent = term ? players.length + '/' + lane.players.length : lane.players.length;
    header.appendChild(count);
    section.appendChild(header);

    const cards = document.createElement('div');
    cards.className = 'cards';
    const sorted = players.slice().sort((a, b) => {
      if (mode === 'name') return a.name.localeCompare(b.name, 'nl');
      const av = a.sd === null ? Infinity : a.sd;
      const bv = b.sd === null ? Infinity : b.sd;
      if (av !== bv) return mode === 'sd-desc' ? bv - av : av - bv;
      return a.name.localeCompare(b.name, 'nl');
    });
    for (const p of sorted) cards.appendChild(makeCard(p));
    if (!sorted.length) {
      const empty = document.createElement('div');
      empty.className = 'empty';
      empty.textContent = 'Geen spelers';
      cards.appendChild(empty);
    }
    section.appendChild(cards);
    if (term && !players.length) section.classList.add('hidden');
    board.appendChild(section);
  }
}

function renderStats() {
  const sds = DATA.players.map(p => p.sd).filter(v => v !== null);
  const best = sds.length ? Math.min.apply(null, sds) : null;
  const without = DATA.lanes.filter(l => l.noTeam).map(l => l.players.length)[0] || 0;
  const items = [
    [DATA.players.length, 'spelers'],
    [DATA.lanes.length, 'teams'],
    [without, 'zonder team'],
    [best === null ? '\u2013' : fmtSd(best), 'beste SD']
  ];
  const stats = document.getElementById('stats');
  for (const item of items) {
    const el = document.createElement('div');
    el.className = 'stat';
    el.innerHTML = '<b>' + item[0] + '</b><span>' + item[1] + '</span>';
    stats.appendChild(el);
  }
}

function renderLegend() {
  const legend = document.getElementById('legend');
  const order = ['Heren', 'Heren 50+', 'Dames', 'Dames 50+', 'C-Mixed', 'Niet gespeeld 2026'];
  for (const cat of order) {
    if (!DATA.lanes.some(l => l.category === cat)) continue;
    const el = document.createElement('span');
    el.innerHTML = '<i style="background:' + CAT_COLORS[cat] + '"></i>' + cat;
    legend.appendChild(el);
  }
}

renderStats();
renderLegend();
render();
q.addEventListener('input', render);
sortSel.addEventListener('change', render);
}

unlockFlow(PAYLOAD, start);
</script>
</body>
</html>
"""

INDELING_TEMPLATE = r"""<!DOCTYPE html>
<html lang="nl">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<meta name="robots" content="noindex, nofollow">
<title>2027 Indeling</title>
<style>__CSS__</style>
</head>
<body>
<header class="top">
  <div class="title">
    <h1>2027 Indeling</h1>
    <p>Sleep kaarten tussen de teams om de nieuwe indeling te maken &middot; dubbelklik op een teamnaam om te hernoemen</p>
  </div>
  <div class="info-group">
    <button type="button" class="info-btn" data-info="sd"><span class="info-icon">i</span>SD</button>
    <button type="button" class="info-btn" data-info="rating"><span class="info-icon">i</span>Rating</button>
  </div>
  <div class="spacer"></div>
  <div class="stats" id="stats"></div>
  <div class="tools">
    <input type="search" id="q" placeholder="Zoek naam&hellip;" autocomplete="off">
    <button class="btn" id="add">+ Nieuw team</button>
    <button class="btn" id="export">Exporteer</button>
    <button class="btn" id="import">Importeer</button>
    <button class="btn ghost" id="reset">Herstel 2026</button>
    <a class="btn" href="__OVERVIEW_HREF__">&larr; Overzicht</a>
    <span class="saved" id="saved">opgeslagen</span>
  </div>
  <input type="file" id="file" accept="application/json,.json" hidden>
</header>
<main class="board" id="board"></main>
__UNLOCK_HTML__
__INFO_HTML__
<script id="payload" type="application/json">__DATA__</script>
<script>__UNLOCK_JS__</script>
<script>__INFO_JS__</script>
<script>
const PAYLOAD = JSON.parse(document.getElementById('payload').textContent);

function start(DATA) {
const PLAYERS = new Map(DATA.players.map(p => [p.id, p]));
const STORAGE_KEY = 'ngf-indeling-2027-v1';

const board = document.getElementById('board');
const q = document.getElementById('q');
const savedEl = document.getElementById('saved');
const fileInput = document.getElementById('file');

const CAT_COLORS = {
  'Heren': '#3b6fd4',
  'Heren 50+': '#c98a1a',
  'Dames': '#d1478a',
  'Dames 50+': '#9b5cc4',
  'C-Mixed': '#2f9e7d',
  'Niet gespeeld 2026': '#8b95a3',
  'Overig': '#6b8cff'
};

const HOLE_COLORS = {
  '36': ['#e3edff', '#1f4fa8', '#bcd3ff'],
  '27': ['#e2f5ee', '#0f7a5a', '#bce8d8'],
  '18': ['#f1e8fb', '#6b33a8', '#ddc9f5'],
  '18/27': ['#e8eefc', '#3d5bbf', '#cbd7f7'],
  '27/36': ['#e0f4f8', '#10708a', '#bfe6ef'],
  'Geen voorkeur': ['#f2f4f7', '#6b7480', '#e4e8ee'],
  '': ['#f2f4f7', '#6b7480', '#e4e8ee']
};

function categoryOf(name) {
  const t = (name || '').trim();
  if (t === 'Niet gespeeld 2026') return 'Niet gespeeld 2026';
  if (t.indexOf('Heren 50+') === 0) return 'Heren 50+';
  if (t.indexOf('Heren') === 0) return 'Heren';
  if (t.indexOf('Dames 50+') === 0) return 'Dames 50+';
  if (t.indexOf('Dames') === 0) return 'Dames';
  if (t.indexOf('C-Mixed') === 0) return 'C-Mixed';
  return 'Overig';
}

function sdScale(sd) {
  const min = 3, max = 36;
  const t = Math.min(1, Math.max(0, (sd - min) / (max - min)));
  const hue = (1 - t) * 128;
  return {
    bg: 'hsl(' + hue + ' 72% 93%)',
    fg: 'hsl(' + hue + ' 55% 30%)',
    bd: 'hsl(' + hue + ' 60% 82%)'
  };
}

function fmtSd(sd) {
  if (sd === null || sd === undefined) return null;
  return sd.toFixed(1).replace('.', ',');
}

function fmtRating(rating) {
  if (rating === null || rating === undefined) return null;
  const v = Math.round(rating * 100) / 100;
  let body = Math.abs(v).toFixed(2).replace(/0+$/, '').replace(/\.$/, '');
  if (body === '') body = '0';
  const sign = v > 0 ? '+' : (v < 0 ? '-' : '');
  return sign + body.replace('.', ',');
}

function ratingColors(rating) {
  const v = Math.round(rating * 100) / 100;
  if (v === 0) return { bg: '#f2f4f7', fg: '#6b7480', bd: '#e4e8ee' };
  const hue = v > 0 ? 140 : 0;
  const mag = Math.min(1, Math.abs(v) / 7);
  const bgL = 94 - mag * 9;
  return {
    bg: 'hsl(' + hue + ' 70% ' + bgL + '%)',
    fg: 'hsl(' + hue + ' 62% ' + (30 - mag * 4) + '%)',
    bd: 'hsl(' + hue + ' 58% ' + (bgL - 9) + '%)'
  };
}

function defaultState() {
  return {
    version: 1,
    lanes: DATA.lanes.map((l, i) => ({
      id: 'lane-' + (i + 1),
      name: l.team,
      players: l.players.map(p => p.id)
    }))
  };
}

function isPoolName(name) {
  return /niet gespeeld|niet ingedeeld|geen team|nog niet/i.test(name || '');
}

function nextLaneId() {
  const ids = new Set(state.lanes.map(l => l.id));
  let n = 1;
  while (ids.has('lane-' + n)) n++;
  return 'lane-' + n;
}

function normalize(raw) {
  if (!raw || !Array.isArray(raw.lanes)) return null;
  const used = new Set();
  const lanes = [];
  for (const entry of raw.lanes) {
    if (!entry) continue;
    const id = entry.id || ('lane-' + Math.random().toString(36).slice(2, 8));
    const name = (entry.name || 'Naamloos').toString();
    const players = [];
    for (const pid of (entry.players || [])) {
      if (PLAYERS.has(pid) && !used.has(pid)) { used.add(pid); players.push(pid); }
    }
    lanes.push({ id: id, name: name, players: players });
  }
  if (!lanes.length) return null;
  const missing = DATA.players.filter(p => !used.has(p.id)).map(p => p.id);
  if (missing.length) {
    let pool = lanes.find(l => isPoolName(l.name));
    if (!pool) { pool = { id: 'lane-pool', name: 'Niet gespeeld 2026', players: [] }; lanes.push(pool); }
    for (const pid of missing) pool.players.push(pid);
  }
  return { version: 1, lanes: lanes };
}

function loadState() {
  try {
    const raw = localStorage.getItem(STORAGE_KEY);
    if (!raw) return null;
    return normalize(JSON.parse(raw));
  } catch (e) {
    return null;
  }
}

let state = loadState() || defaultState();
let savedTimer = null;

function save() {
  try { localStorage.setItem(STORAGE_KEY, JSON.stringify(state)); } catch (e) { /* ignore */ }
  savedEl.classList.add('show');
  clearTimeout(savedTimer);
  savedTimer = setTimeout(() => savedEl.classList.remove('show'), 1100);
}

function applyAccent(section, name) {
  section.style.setProperty('--accent', CAT_COLORS[categoryOf(name)] || '#6b8cff');
}

function makeCard(p, laneId) {
  const card = document.createElement('div');
  card.className = 'card';
  card.draggable = true;
  card.dataset.id = p.id;
  card.dataset.name = p.name.toLowerCase();
  const hc = HOLE_COLORS[p.holes] || HOLE_COLORS[''];
  card.style.setProperty('--hole', hc[1]);
  card.style.setProperty('--hole-bg', hc[0]);
  card.style.setProperty('--hole-fg', hc[1]);
  card.style.setProperty('--hole-bd', hc[2]);

  const name = document.createElement('div');
  name.className = 'name';
  name.textContent = p.name;
  card.appendChild(name);

  const meta = document.createElement('div');
  meta.className = 'meta';

  const holes = document.createElement('span');
  holes.className = 'badge holes';
  holes.textContent = p.holes ? p.holes + ' holes' : 'holes onbekend';
  meta.appendChild(holes);

  const sd = document.createElement('span');
  const val = fmtSd(p.sd);
  if (val === null) {
    sd.className = 'badge sd none';
    sd.textContent = 'SD n.v.t.';
  } else {
    sd.className = 'badge sd';
    const c = sdScale(p.sd);
    sd.style.setProperty('--sd-bg', c.bg);
    sd.style.setProperty('--sd-fg', c.fg);
    sd.style.setProperty('--sd-bd', c.bd);
    sd.textContent = 'SD ' + val;
    sd.title = 'Gemiddeld SD ' + p.sd;
  }
  meta.appendChild(sd);

  const rt = document.createElement('span');
  const rval = fmtRating(p.rating);
  if (rval === null) {
    rt.className = 'badge sd none';
    rt.textContent = 'Rating n.v.t.';
  } else {
    rt.className = 'badge sd';
    const rc = ratingColors(p.rating);
    rt.style.setProperty('--sd-bg', rc.bg);
    rt.style.setProperty('--sd-fg', rc.fg);
    rt.style.setProperty('--sd-bd', rc.bd);
    rt.textContent = 'Rating ' + rval;
    rt.title = 'Rating 2026 (Rt) ' + p.rating;
  }
  meta.appendChild(rt);

  card.appendChild(meta);

  card.addEventListener('dragstart', ev => {
    dragged = { id: p.id, from: laneId };
    card.classList.add('dragging');
    placeholder.style.height = card.offsetHeight + 'px';
    ev.dataTransfer.effectAllowed = 'move';
    try { ev.dataTransfer.setData('text/plain', p.id); } catch (e) { /* ignore */ }
  });
  card.addEventListener('dragend', cleanupDrag);
  return card;
}

let dragged = null;
const placeholder = document.createElement('div');
placeholder.className = 'placeholder';

function getDragAfterElement(container, y) {
  const els = Array.prototype.slice.call(container.querySelectorAll('.card:not(.dragging)'));
  let closest = { offset: Number.NEGATIVE_INFINITY, element: null };
  for (const child of els) {
    const box = child.getBoundingClientRect();
    const offset = y - box.top - box.height / 2;
    if (offset < 0 && offset > closest.offset) closest = { offset: offset, element: child };
  }
  return closest.element;
}

function render() {
  board.innerHTML = '';
  const term = q.value.trim().toLowerCase();
  let assigned = 0;

  for (const lane of state.lanes) {
    const section = document.createElement('section');
    section.className = 'lane';
    section.dataset.laneId = lane.id;
    applyAccent(section, lane.name);

    const header = document.createElement('header');
    const h2 = document.createElement('h2');
    h2.textContent = lane.name;
    h2.contentEditable = 'true';
    h2.spellcheck = false;
    h2.addEventListener('keydown', ev => {
      if (ev.key === 'Enter') { ev.preventDefault(); h2.blur(); }
      else if (ev.key === 'Escape') { h2.textContent = lane.name; h2.blur(); }
    });
    h2.addEventListener('blur', () => {
      const clean = h2.textContent.replace(/\s+/g, ' ').trim();
      if (clean && clean !== lane.name) { lane.name = clean; save(); }
      h2.textContent = lane.name;
      applyAccent(section, lane.name);
    });
    header.appendChild(h2);

    const count = document.createElement('span');
    count.className = 'count';
    count.textContent = lane.players.length;
    header.appendChild(count);

    const del = document.createElement('button');
    del.className = 'del';
    del.textContent = '\u00d7';
    del.title = 'Team verwijderen';
    del.addEventListener('click', () => deleteLane(lane.id));
    header.appendChild(del);

    section.appendChild(header);

    const cards = document.createElement('div');
    cards.className = 'cards';
    for (const pid of lane.players) {
      const p = PLAYERS.get(pid);
      if (!p) continue;
      const card = makeCard(p, lane.id);
      if (term && p.name.toLowerCase().indexOf(term) === -1) card.classList.add('dimmed');
      cards.appendChild(card);
      assigned++;
    }
    if (!lane.players.length) {
      const empty = document.createElement('div');
      empty.className = 'empty';
      empty.textContent = 'Sleep hier een kaart';
      cards.appendChild(empty);
    }
    section.appendChild(cards);
    board.appendChild(section);
  }

  renderStats(assigned);
}

function renderStats(assigned) {
  const items = [
    [DATA.players.length, 'spelers'],
    [state.lanes.length, 'teams'],
    [assigned, 'ingedeeld']
  ];
  const stats = document.getElementById('stats');
  stats.innerHTML = '';
  for (const item of items) {
    const el = document.createElement('div');
    el.className = 'stat';
    el.innerHTML = '<b>' + item[0] + '</b><span>' + item[1] + '</span>';
    stats.appendChild(el);
  }
}

function cleanupDrag() {
  const card = board.querySelector('.card.dragging');
  if (card) card.classList.remove('dragging');
  if (placeholder.parentNode) placeholder.parentNode.removeChild(placeholder);
  const target = board.querySelector('.lane.drop-target');
  if (target) target.classList.remove('drop-target');
  dragged = null;
}

board.addEventListener('dragover', ev => {
  if (!dragged) return;
  ev.preventDefault();
  ev.dataTransfer.dropEffect = 'move';
  const laneEl = ev.target.closest('.lane');
  if (!laneEl) return;
  const cards = laneEl.querySelector('.cards');
  const after = getDragAfterElement(cards, ev.clientY);
  if (after === null) cards.appendChild(placeholder);
  else cards.insertBefore(placeholder, after);
  for (const l of board.querySelectorAll('.lane.drop-target')) l.classList.remove('drop-target');
  laneEl.classList.add('drop-target');
});

board.addEventListener('drop', ev => {
  if (!dragged) return;
  ev.preventDefault();
  const laneEl = ev.target.closest('.lane');
  if (!laneEl || !placeholder.parentNode) { cleanupDrag(); render(); return; }

  const cards = laneEl.querySelector('.cards');
  let index = 0;
  for (const child of cards.children) {
    if (child === placeholder) break;
    if (child.classList.contains('card') && !child.classList.contains('dragging')) index++;
  }

  const fromLane = state.lanes.find(l => l.id === dragged.from);
  const toLane = state.lanes.find(l => l.id === laneEl.dataset.laneId);
  if (!fromLane || !toLane) { cleanupDrag(); render(); return; }

  const fromIndex = fromLane.players.indexOf(dragged.id);
  if (fromIndex === -1) { cleanupDrag(); render(); return; }
  fromLane.players.splice(fromIndex, 1);
  toLane.players.splice(index, 0, dragged.id);

  cleanupDrag();
  save();
  render();
});

function deleteLane(id) {
  if (state.lanes.length <= 1) return;
  const lane = state.lanes.find(l => l.id === id);
  if (!lane) return;
  if (lane.players.length) {
    if (!confirm('"' + lane.name + '" bevat ' + lane.players.length + ' speler(s).\nVerplaatsen naar "Niet gespeeld 2026"?')) return;
    let pool = state.lanes.find(l => l.id !== id && isPoolName(l.name));
    if (!pool) {
      pool = { id: nextLaneId(), name: 'Niet gespeeld 2026', players: [] };
      state.lanes.push(pool);
    }
    for (const pid of lane.players) pool.players.push(pid);
  }
  state.lanes = state.lanes.filter(l => l.id !== id);
  save();
  render();
}

function addLane() {
  const lane = { id: nextLaneId(), name: 'Nieuw team', players: [] };
  state.lanes.push(lane);
  save();
  render();
  const h2 = board.querySelector('.lane[data-lane-id="' + lane.id + '"] h2');
  if (h2) {
    h2.focus();
    const range = document.createRange();
    range.selectNodeContents(h2);
    const sel = window.getSelection();
    sel.removeAllRanges();
    sel.addRange(range);
  }
}

document.getElementById('add').addEventListener('click', addLane);

document.getElementById('reset').addEventListener('click', () => {
  if (!confirm('De huidige indeling wordt vervangen door de teams van 2026. Doorgaan?')) return;
  state = defaultState();
  save();
  render();
});

document.getElementById('export').addEventListener('click', () => {
  const blob = new Blob([JSON.stringify(state, null, 2)], { type: 'application/json' });
  const url = URL.createObjectURL(blob);
  const a = document.createElement('a');
  a.href = url;
  a.download = '2027-indeling.json';
  document.body.appendChild(a);
  a.click();
  document.body.removeChild(a);
  URL.revokeObjectURL(url);
});

document.getElementById('import').addEventListener('click', () => fileInput.click());

fileInput.addEventListener('change', () => {
  const file = fileInput.files && fileInput.files[0];
  if (!file) return;
  const reader = new FileReader();
  reader.onload = () => {
    try {
      const next = normalize(JSON.parse(reader.result));
      if (!next) throw new Error('ongeldig bestand');
      state = next;
      save();
      render();
    } catch (e) {
      alert('Kon dit bestand niet lezen: ' + e.message);
    }
    fileInput.value = '';
  };
  reader.readAsText(file);
});

q.addEventListener('input', render);
render();
}

unlockFlow(PAYLOAD, start);
</script>
</body>
</html>
"""


def node_path() -> str:
    node = shutil.which("node")
    if not node:
        raise SystemExit(
            "Node.js is nodig om te versleutelen (alleen tijdens het bouwen). "
            "Installeer Node of bouw zonder --encrypt.")
    return node


def encrypt_payload(plaintext: str, password: str) -> str:
    """Encrypt the JSON payload with PBKDF2-SHA256 + AES-256-GCM via Node."""
    env = dict(os.environ, TC_PASSWORD=password, TC_ITERATIONS=str(PBKDF2_ITERATIONS))
    proc = subprocess.run(
        [node_path(), str(ENCRYPTOR)],
        input=plaintext.encode("utf-8"),
        stdout=subprocess.PIPE, stderr=subprocess.PIPE, env=env)
    if proc.returncode != 0:
        raise SystemExit("Versleutelen mislukt: " + proc.stderr.decode("utf-8", "replace"))
    return proc.stdout.decode("utf-8")


def read_password(args) -> str:
    if args.password:
        return args.password
    first = getpass.getpass("Wachtwoord voor de pagina's: ")
    if not first:
        raise SystemExit("Leeg wachtwoord is niet toegestaan.")
    if len(first) < 8:
        raise SystemExit("Kies een wachtwoord van minstens 8 tekens.")
    second = getpass.getpass("Bevestig het wachtwoord: ")
    if first != second:
        raise SystemExit("De wachtwoorden komen niet overeen.")
    return first


def render_pages(payload_script: str, docs: bool) -> None:
    if docs:
        DOCS_DIR.mkdir(exist_ok=True)
        overview_out = DOCS_DIR / "index.html"
        indeling_out = DOCS_DIR / "indeling.html"
        overview_href = "index.html"
        indeling_href = "indeling.html"
        (DOCS_DIR / ".nojekyll").write_text("", encoding="utf-8")
    else:
        overview_out = ROOT / OVERVIEW_NAME
        indeling_out = ROOT / INDELING_NAME
        overview_href = OVERVIEW_NAME.replace(" ", "%20")
        indeling_href = INDELING_NAME.replace(" ", "%20")

    common = {
        "__DATA__": payload_script,
        "__UNLOCK_HTML__": UNLOCK_HTML,
        "__UNLOCK_JS__": UNLOCK_JS,
        "__INFO_HTML__": INFO_HTML,
        "__INFO_JS__": INFO_JS,
    }

    overview = OVERVIEW_TEMPLATE.replace("__CSS__", BASE_CSS)
    indeling = INDELING_TEMPLATE.replace("__CSS__", BASE_CSS + INDELING_CSS)
    overview = overview.replace("__INDELING_HREF__", indeling_href)
    indeling = indeling.replace("__OVERVIEW_HREF__", overview_href)
    for token, value in common.items():
        overview = overview.replace(token, value)
        indeling = indeling.replace(token, value)

    overview_out.write_text(overview, encoding="utf-8")
    indeling_out.write_text(indeling, encoding="utf-8")

    print(f"  -> {overview_out}")
    print(f"  -> {indeling_out}")


def main():
    parser = argparse.ArgumentParser(description="Bouw de TC Ranking pagina's.")
    parser.add_argument("--encrypt", action="store_true",
                        help="versleutel de spelersgegevens met een wachtwoord")
    parser.add_argument("--password", default=None,
                        help="wachtwoord (niet-interactief; liever zonder deze optie)")
    parser.add_argument("--docs", action="store_true",
                        help="schrijf docs/index.html en docs/indeling.html voor GitHub Pages")
    parser.add_argument("--force-plain-docs", action="store_true",
                        help="sta een NIET-versleutelde docs-build toe (niet voor publicatie)")
    args = parser.parse_args()

    if args.docs and not (args.encrypt or args.password) and not args.force_plain_docs:
        raise SystemExit(
            "Weiger een niet-versleutelde docs-build te maken: GitHub Pages is openbaar.\n"
            "Gebruik --encrypt, of --force-plain-docs voor lokaal testen.")

    ratings = load_ratings()
    players = load_players(ratings)
    lanes = build_lanes(players)
    data = {"players": players, "lanes": lanes}
    plaintext = json.dumps(data, ensure_ascii=False)

    if args.encrypt or args.password:
        password = read_password(args)
        encrypted = json.loads(encrypt_payload(plaintext, password))
        payload_script = json.dumps({"enc": encrypted}, ensure_ascii=False)
        print(f"{len(players)} spelers in {len(lanes)} teams (versleuteld)")
    else:
        payload_script = json.dumps({"plain": data}, ensure_ascii=False)
        print(f"{len(players)} spelers in {len(lanes)} teams (niet versleuteld)")

    render_pages(payload_script, args.docs)


if __name__ == "__main__":
    main()
